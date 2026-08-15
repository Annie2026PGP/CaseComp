"""GRAD 3.0 -- Galderma Sunscreen Challenge. Builds the 2-slide executive
summary version of the full strategy deck (build_deck.py).

Run: python3 cases/2026-08-galderma-grad3-sunscreen/deck/build_deck_2slide.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = "cases/2026-08-galderma-grad3-sunscreen/deck"
ASSETS = f"{HERE}/assets"
OUT = f"{HERE}/../submission/GRAD3_Galderma_Sunscreen_2Slide_Summary.pptx"

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xFB, 0xF7, 0xF0)
INK = RGBColor(0x1E, 0x26, 0x26)
TEAL = RGBColor(0x0B, 0x5F, 0x63)
TEAL_DARK = RGBColor(0x08, 0x3F, 0x42)
TEAL_LIGHT = RGBColor(0xBF, 0xDC, 0xDA)
CORAL = RGBColor(0xE8, 0x73, 0x4A)
CORAL_LIGHT = RGBColor(0xF6, 0xC7, 0xB4)
GOLD = RGBColor(0xE3, 0xA9, 0x3E)
GREY = RGBColor(0x8A, 0x94, 0x94)
GREY_TEXT = RGBColor(0x55, 0x5F, 0x5F)

FONT_HEAD = "Georgia"
FONT = "Arial"
MIN_PT = 10

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SLIDE_N = 0


def add_slide(bg=PAPER):
    global SLIDE_N
    SLIDE_N += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def _apply(p, rd):
    r = p.add_run()
    r.text = rd["text"]
    size = rd.get("size", 12)
    assert size >= MIN_PT, f"font too small: {size}"
    r.font.size = Pt(size)
    r.font.color.rgb = rd.get("color", INK)
    r.font.bold = rd.get("bold", False)
    r.font.italic = rd.get("italic", False)
    r.font.name = rd.get("font", FONT)


def textbox(slide, l, t, w, h, text, size=14, color=INK, bold=False, italic=False,
            align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
            space_after=0):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        _apply(p, {"text": line, "size": size, "color": color, "bold": bold,
                   "italic": italic, "font": font})
    return box


def rich(slide, l, t, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.2, space_after=6):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        for rd in runs:
            _apply(p, rd)
    return box


def rect(slide, l, t, w, h, fill=TEAL, line_color=None, line_w=1.0, radius=None):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(st, Inches(l), Inches(t), Inches(w), Inches(h))
    if radius is not None:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    return sh


def picture(slide, path, l, t, w=None, h=None):
    kwargs = {}
    if w is not None:
        kwargs["width"] = Inches(w)
    if h is not None:
        kwargs["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(l), Inches(t), **kwargs)


def header(slide, kicker, title, num):
    n_lines = title.count("\n") + 1
    title_h = 0.36 * n_lines
    textbox(slide, 0.6, 0.3, 10, 0.3, kicker, size=11.5, color=CORAL, bold=True, font=FONT)
    textbox(slide, 0.6, 0.58, 11.8, title_h + 0.1, title, size=21, color=TEAL_DARK, bold=True,
            font=FONT_HEAD, line_spacing=1.05)
    rect(slide, 0.6, 0.58 + title_h + 0.12, 0.9, 0.04, fill=CORAL)
    footer(slide, num)


def footer(slide, num):
    textbox(slide, 0.6, 7.16, 8, 0.26, "GRAD 3.0 -- Galderma Sunscreen Challenge -- Executive Summary",
            size=10, color=GREY)
    textbox(slide, 11.9, 7.16, 0.9, 0.26, str(num), size=10, color=GREY, align=PP_ALIGN.RIGHT)


def stat_card(slide, l, t, w, h, big, label, color=TEAL):
    rect(slide, l, t, w, h, fill=WHITE, line_color=color, line_w=1.3, radius=0.09)
    textbox(slide, l + 0.12, t + 0.1, w - 0.24, h * 0.52, big, size=18, color=color, bold=True,
            font=FONT_HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, l + 0.12, t + h * 0.54, w - 0.24, h * 0.42, label, size=10, color=GREY_TEXT,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, line_spacing=1.15)


# ===========================================================================
# SLIDE 1 -- THE OPPORTUNITY
# ===========================================================================
s = add_slide()
header(s, "GRAD 3.0 -- GALDERMA SUNSCREEN CHALLENGE",
       "Galderma has the derm credibility.\nIt doesn't have the format range -- yet.", 1)

rich(s, 0.6, 1.55, 6.05, 2.15, [
    [{"text": "The gap is translation, not science: ", "size": 12.5, "bold": True, "color": TEAL_DARK},
     {"text": "Cetaphil and Biluma already carry the dermatological credibility indie D2C "
              "brands (Minimalist, Dot & Key, The Derma Co.) are still building. What they "
              "haven't built is the format range and content presence that makes that "
              "credibility legible to India's 25-35 sunscreen buyer.", "size": 12.5}],
    [{"text": "Three consumer jobs -- Commuter, Skintok Regular, Occasional Buyer -- "
              "share one blocker: ", "size": 12.5, "bold": True, "color": TEAL_DARK},
     {"text": "white cast, the single most cited reason Indian consumers skip reapplication.",
      "size": 12.5}],
], line_spacing=1.32, space_after=12)

stat_card(s, 0.6, 3.9, 1.95, 1.15, "6-9%", "India sun care\nCAGR (2024-30)", color=TEAL)
stat_card(s, 2.68, 3.9, 1.95, 1.15, "10.8%", "sunscreen stick\nCAGR (fastest)", color=CORAL)
stat_card(s, 4.76, 3.9, 1.89, 1.15, "White cast", "#1 reapplication\nbarrier, India", color=GOLD)

rect(s, 0.6, 5.2, 6.05, 1.55, fill=TEAL_DARK, radius=0.08)
rich(s, 0.85, 5.36, 5.55, 1.25, [[
    {"text": "The white space: ", "size": 12, "bold": True, "color": GOLD},
    {"text": "high derm credibility + high format/sensorial modernity -- where La "
             "Roche-Posay Anthelios plays internationally but hasn't fully localized for "
             "Indian skin and climate. Nobody owns it here yet.", "size": 12, "color": WHITE},
]], line_spacing=1.3)

picture(s, f"{ASSETS}/competitive_map.png", 6.85, 1.7, h=4.5)

# ===========================================================================
# SLIDE 2 -- THE STRATEGY
# ===========================================================================
s = add_slide()
header(s, "GRAD 3.0 -- GALDERMA SUNSCREEN CHALLENGE",
       "Two brands, two systems, one standard", 2)

textbox(s, 0.6, 1.5, 12.1, 0.4,
        "A 3-5 year portfolio, built to eliminate white cast portfolio-wide and win the "
        "trial-to-repeat path via quick commerce and dermatologist-creator partnerships.",
        size=12, color=GREY_TEXT, italic=True, line_spacing=1.25)

picture(s, f"{ASSETS}/portfolio_hierarchy.png", 4.2, 2.0, h=2.9)

roadmap = [
    ("PHASE 1 -- Year 1", "Fix the core", "Zero-white-cast reformulation portfolio-wide; "
     "launch Cetaphil Sun Stick + Biluma Sun Mist", TEAL),
    ("PHASE 2 -- Yrs 2-3", "Build the premium tier", "Launch Biluma Sun Tinted + Cetaphil "
     "Gel-Cream; scale dermatologist-creator content", CORAL),
    ("PHASE 3 -- Yrs 3-5", "Defend and extend", "Trial-size SKUs across the range; "
     "quick-commerce reorder mechanics; evaluate Kids/Family", GOLD),
]
x = 0.6
w = 3.95
for tag, title, body, color in roadmap:
    rect(s, x, 5.05, w, 0.42, fill=color)
    textbox(s, x + 0.14, 5.05, w - 0.28, 0.42, tag, size=10.5, color=WHITE, bold=True,
            anchor=MSO_ANCHOR.MIDDLE, font=FONT)
    rect(s, x, 5.47, w, 1.3, fill=WHITE, line_color=color, line_w=1.1)
    rich(s, x + 0.15, 5.58, w - 0.3, 1.1, [
        [{"text": title, "size": 11.5, "bold": True, "color": TEAL_DARK}],
        [{"text": body, "size": 10}],
    ], line_spacing=1.25, space_after=4)
    x += w + 0.19

rect(s, 0.6, 6.85, 12.13, 0.02, fill=GREY)
textbox(s, 0.6, 6.89, 12.13, 0.24,
        "Full detail -- SPF architecture, claims, pricing ladder, GTM by channel -- in "
        "GRAD3_Galderma_Sunscreen_Portfolio_Strategy.pptx (10-slide version).",
        size=10, color=GREY, italic=True)

prs.save(OUT)
print("Saved:", OUT)
print("Total slides:", SLIDE_N)
