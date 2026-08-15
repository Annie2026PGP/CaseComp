"""GRAD 3.0 -- Galderma Sunscreen Challenge. Builds the strategy deck.

Run: python3 cases/2026-08-galderma-grad3-sunscreen/deck/build_deck.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = "cases/2026-08-galderma-grad3-sunscreen/deck"
ASSETS = f"{HERE}/assets"
OUT = f"{HERE}/../submission/GRAD3_Galderma_Sunscreen_Portfolio_Strategy.pptx"

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xFB, 0xF7, 0xF0)
INK = RGBColor(0x1E, 0x26, 0x26)
TEAL = RGBColor(0x0B, 0x5F, 0x63)
TEAL_DARK = RGBColor(0x08, 0x3F, 0x42)
TEAL_LIGHT = RGBColor(0xBF, 0xDC, 0xDA)
CORAL = RGBColor(0xE8, 0x73, 0x4A)
CORAL_LIGHT = RGBColor(0xF6, 0xC7, 0xB4)
GOLD = RGBColor(0xE3, 0xA9, 0x3E)
GREY = RGBColor(0x8A, 0x94, 0x94)
GREY_TEXT = RGBColor(0x55, 0x5F, 0x5F)

FONT_HEAD = "Georgia"
FONT = "Arial"
MIN_PT = 10

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SLIDE_N = 0


def add_slide(bg=PAPER):
    global SLIDE_N
    SLIDE_N += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def _apply(p, rd):
    r = p.add_run()
    r.text = rd["text"]
    size = rd.get("size", 12)
    assert size >= MIN_PT, f"font too small: {size}"
    r.font.size = Pt(size)
    r.font.color.rgb = rd.get("color", INK)
    r.font.bold = rd.get("bold", False)
    r.font.italic = rd.get("italic", False)
    r.font.name = rd.get("font", FONT)


def textbox(slide, l, t, w, h, text, size=14, color=INK, bold=False, italic=False,
            align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
            space_after=0):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        _apply(p, {"text": line, "size": size, "color": color, "bold": bold,
                   "italic": italic, "font": font})
    return box


def rich(slide, l, t, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.2, space_after=6):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        for rd in runs:
            _apply(p, rd)
    return box


def rect(slide, l, t, w, h, fill=TEAL, line_color=None, line_w=1.0, radius=None):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(st, Inches(l), Inches(t), Inches(w), Inches(h))
    if radius is not None:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    return sh


def picture(slide, path, l, t, w=None, h=None):
    kwargs = {}
    if w is not None:
        kwargs["width"] = Inches(w)
    if h is not None:
        kwargs["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(l), Inches(t), **kwargs)


def header(slide, kicker, title, num):
    n_lines = title.count("\n") + 1
    title_h = 0.42 * n_lines
    textbox(slide, 0.6, 0.35, 10, 0.32, kicker, size=11.5, color=CORAL, bold=True, font=FONT)
    textbox(slide, 0.6, 0.64, 11.5, title_h + 0.1, title, size=24, color=TEAL_DARK, bold=True,
            font=FONT_HEAD, line_spacing=1.05)
    rect(slide, 0.6, 0.64 + title_h + 0.14, 0.9, 0.045, fill=CORAL)
    footer(slide, num)


def footer(slide, num):
    textbox(slide, 0.6, 7.14, 8, 0.28, "GRAD 3.0 -- Galderma Sunscreen Challenge", size=10, color=GREY)
    textbox(slide, 11.9, 7.14, 0.9, 0.28, str(num), size=10, color=GREY, align=PP_ALIGN.RIGHT)


def stat_card(slide, l, t, w, h, big, label, color=TEAL):
    rect(slide, l, t, w, h, fill=WHITE, line_color=color, line_w=1.4, radius=0.08)
    textbox(slide, l + 0.18, t + 0.14, w - 0.36, h * 0.55, big, size=22, color=color, bold=True,
            font=FONT_HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, l + 0.18, t + h * 0.58, w - 0.36, h * 0.4, label, size=10.5, color=GREY_TEXT,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, line_spacing=1.2)


# ===========================================================================
# COVER
# ===========================================================================
s = add_slide(bg=TEAL_DARK)
rect(s, 0, 0, 13.333, 7.5, fill=TEAL_DARK)
rect(s, 0, 6.85, 13.333, 0.06, fill=CORAL)
textbox(s, 0.9, 1.1, 6, 0.35, "GRAD 3.0 -- GALDERMA SUNSCREEN CHALLENGE", size=13,
        color=GOLD, bold=True, font=FONT)
textbox(s, 0.9, 1.75, 11, 2.0, "Two Brands, One System:\nA Sunscreen Portfolio for\nIndia's Next Generation",
        size=40, color=WHITE, bold=True, font=FONT_HEAD, line_spacing=1.08)
textbox(s, 0.9, 4.05, 10.2, 0.9,
        "A 3-5 year portfolio and go-to-market strategy for Cetaphil Sun and Biluma Sun, "
        "built to convert Galderma's dermatological credibility into the formats, claims "
        "and channels India's 25-35 sunscreen consumer already trusts.",
        size=14.5, color=RGBColor(0xD8, 0xE6, 0xE4), line_spacing=1.4, font=FONT)
rect(s, 0.92, 5.35, 3.2, 0.02, fill=RGBColor(0x3E, 0x7A, 0x7D))
textbox(s, 0.9, 5.55, 8, 0.3, "Strategic consultants to Galderma India | GRAD 3.0 submission",
        size=11, color=GREY, font=FONT)
rect(s, 0.9, 6.0, 8.5, 0.62, fill=RGBColor(0x0F, 0x4D, 0x50), radius=0.15)
textbox(s, 1.12, 6.14, 8.1, 0.4,
        "Secondary-research-driven strategy case -- no dataset or exhibits were provided; "
        "all market figures cited in research/sources.md.",
        size=10.5, color=RGBColor(0xB8, 0xCE, 0xCC), italic=True, font=FONT, line_spacing=1.25)

# ===========================================================================
# SLIDE: The problem, sharpened
# ===========================================================================
s = add_slide()
header(s, "THE PROBLEM", "Galderma has the derm credibility.\nIt doesn't yet have the format range.", 2)

rich(s, 0.6, 1.75, 5.9, 4.9, [
    [{"text": "Cetaphil Sun and Biluma Sun deliver trusted, dermatologically tested protection. "
              "But the 25-35 Indian consumer is choosing sunscreen on format, sensorial feel and "
              "content presence as much as on protection science.", "size": 13.5}],
    [{"text": "That's a translation gap, not a science gap: ", "size": 13.5, "bold": True, "color": TEAL_DARK},
     {"text": "Galderma already owns the credibility indie D2C brands are still building. What's "
              "missing is the format range and claims architecture to make that credibility "
              "legible to a skincare-literate buyer.", "size": 13.5}],
], line_spacing=1.4, space_after=16)

stat_card(s, 0.6, 4.5, 2.75, 1.55, "6-9%", "India sun care CAGR\n(2024-2030, est.)", color=TEAL)
stat_card(s, 3.5, 4.5, 2.75, 1.55, "10.8%", "sunscreen stick format\nCAGR -- fastest-growing", color=CORAL)
stat_card(s, 0.6, 6.15, 5.65, 0.85, "White cast", "the #1 cited reason for skipped reapplication\namong Indian consumers", color=GOLD)

picture(s, f"{ASSETS}/competitive_map.png", 6.75, 1.7, w=6.0)

# ===========================================================================
# SLIDE: The consumer -- three jobs, one blocker
# ===========================================================================
s = add_slide()
header(s, "CONSUMER DEEP DIVE", "Three jobs, one shared blocker", 3)

personas = [
    ("THE COMMUTER", "Daily, defensive habit before stepping out. Wants fast absorption and "
     "zero mid-day retouching.", "Reapplication friction -- white cast + greasy feel mean the "
     "2pm top-up just doesn't happen.", TEAL),
    ("THE SKINTOK REGULAR", "Sunscreen as step 5-7 of a routine, filmed or content-influenced. "
     "Wants ingredient literacy and a pack worth showing.", "Today's format reads \"medicinal,\" "
     "not \"routine-worthy\" -- credibility isn't signaled visually.", CORAL),
    ("THE OCCASIONAL BUYER", "Reaches for sunscreen only for high-exposure occasions -- beach, "
     "wedding, a long outdoor day.", "Low category education -- buys on brand recall and pack "
     "size, not SPF/PA understanding.", GOLD),
]
x = 0.6
w = 3.95
for name, job, barrier, color in personas:
    rect(s, x, 1.75, w, 0.6, fill=color)
    textbox(s, x + 0.16, 1.75, w - 0.32, 0.6, name, size=13, color=WHITE, bold=True,
            anchor=MSO_ANCHOR.MIDDLE, font=FONT)
    rect(s, x, 2.35, w, 3.05, fill=WHITE, line_color=color, line_w=1.2)
    rich(s, x + 0.2, 2.55, w - 0.4, 1.4, [
        [{"text": "JOB: ", "size": 10.5, "bold": True, "color": color},
         {"text": job, "size": 10.5}],
    ], line_spacing=1.3)
    rich(s, x + 0.2, 4.05, w - 0.4, 1.2, [
        [{"text": "BARRIER: ", "size": 10.5, "bold": True, "color": color},
         {"text": barrier, "size": 10.5}],
    ], line_spacing=1.3)
    x += w + 0.19

rect(s, 0.6, 5.6, 12.13, 1.15, fill=TEAL_DARK, radius=0.08)
rich(s, 0.85, 5.78, 11.6, 0.85, [[
    {"text": "The unifying insight: ", "size": 12.5, "bold": True, "color": GOLD},
    {"text": "white cast is not a niche complaint -- it hits every persona differently, and fixing "
             "it across formats is the single highest-leverage move available to this portfolio.",
     "size": 12.5, "color": WHITE},
]], line_spacing=1.35)

# ===========================================================================
# SLIDE: Competitive landscape (full map)
# ===========================================================================
s = add_slide()
header(s, "COMPETITIVE LANDSCAPE", "The white space: derm credibility, modern format", 4)
picture(s, f"{ASSETS}/competitive_map.png", 3.55, 1.55, h=5.35)
rich(s, 0.6, 1.75, 2.75, 5.0, [
    [{"text": "Indie-derma brands", "size": 12, "bold": True, "color": TEAL}],
    [{"text": "have won the format and content narrative with far less R&D pedigree than "
              "Galderma actually has.", "size": 10.8}],
    [{"text": "La Roche-Posay Anthelios", "size": 12, "bold": True, "color": TEAL_DARK}],
    [{"text": "is the international benchmark for this exact move -- derm heritage, "
              "multi-format range -- not yet fully localized for India.", "size": 10.8}],
    [{"text": "Cetaphil / Biluma today", "size": 12, "bold": True, "color": CORAL}],
    [{"text": "sit bottom-right: real credibility, narrow format range. That gap is the "
              "strategy.", "size": 10.8}],
], line_spacing=1.3, space_after=10)

# ===========================================================================
# SLIDE: Portfolio strategy overview
# ===========================================================================
s = add_slide()
header(s, "PORTFOLIO STRATEGY", "Two brands, two systems, one standard", 5)
textbox(s, 0.6, 1.5, 11.8, 0.55,
        "Cetaphil and Biluma stop being two sunscreens and become two systems -- each built "
        "around a distinct job, both engineered to eliminate white cast.",
        size=13, color=GREY_TEXT, italic=True, line_spacing=1.3)
picture(s, f"{ASSETS}/portfolio_hierarchy.png", 1.62, 2.15, h=4.85)

# ===========================================================================
# SLIDE: SPF architecture, claims, ingredients
# ===========================================================================
s = add_slide()
header(s, "PORTFOLIO STRATEGY", "SPF architecture, claims and ingredients", 6)

rect(s, 0.6, 1.65, 12.13, 1.1, fill=GOLD, radius=0.06)
rich(s, 0.85, 1.83, 11.6, 0.8, [[
    {"text": "SPF50+ PA++++ across the entire portfolio. ", "size": 13, "bold": True, "color": TEAL_DARK},
    {"text": "India's climate makes a lower-SPF \"everyday\" tier a false economy. Differentiate "
             "by format and finish, not protection level -- every Galderma sunscreen is \"the good "
             "one,\" which also avoids fragmenting a still low-education category.",
     "size": 12.5, "color": TEAL_DARK},
]], line_spacing=1.32)

colw = 5.85
rect(s, 0.6, 3.0, colw, 3.6, fill=WHITE, line_color=TEAL, line_w=1.4, radius=0.06)
rect(s, 0.6, 3.0, colw, 0.6, fill=TEAL)
textbox(s, 0.8, 3.0, colw - 0.4, 0.6, "CETAPHIL SUN -- claims & ingredients", size=12.5,
        color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
rich(s, 0.85, 3.78, colw - 0.5, 2.7, [
    [{"text": "Portfolio-wide baseline: ", "size": 11.5, "bold": True, "color": TEAL_DARK},
     {"text": "\"No White Cast\" -- extended from Biluma-only to every SKU.", "size": 11.5}],
    [{"text": "Lead claim: ", "size": 11.5, "bold": True, "color": TEAL_DARK},
     {"text": "\"Dermatologist-recommended, all-day comfort.\"", "size": 11.5}],
    [{"text": "Retained: ", "size": 11.5, "bold": True, "color": TEAL_DARK},
     {"text": "Vitamin E -- Cetaphil's existing sensitive-skin anchor.", "size": 11.5}],
    [{"text": "New: ", "size": 11.5, "bold": True, "color": TEAL_DARK},
     {"text": "lightweight/non-greasy claim on the stick and gel-cream formats.", "size": 11.5}],
], line_spacing=1.35, space_after=12)

rect(s, 6.7, 3.0, colw, 3.6, fill=WHITE, line_color=CORAL, line_w=1.4, radius=0.06)
rect(s, 6.7, 3.0, colw, 0.6, fill=CORAL)
textbox(s, 6.9, 3.0, colw - 0.4, 0.6, "BILUMA SUN -- claims & ingredients", size=12.5,
        color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
rich(s, 6.95, 3.78, colw - 0.5, 2.7, [
    [{"text": "Portfolio-wide baseline: ", "size": 11.5, "bold": True, "color": RGBColor(0xA8, 0x43, 0x1E)},
     {"text": "\"No White Cast\" -- Biluma's existing claim, now the standard.", "size": 11.5}],
    [{"text": "Lead claim: ", "size": 11.5, "bold": True, "color": RGBColor(0xA8, 0x43, 0x1E)},
     {"text": "\"Protection + Visible Skin Goal.\"", "size": 11.5}],
    [{"text": "Retained: ", "size": 11.5, "bold": True, "color": RGBColor(0xA8, 0x43, 0x1E)},
     {"text": "Lactic Acid + Hyaluronic Acid -- Biluma's existing actives.", "size": 11.5}],
    [{"text": "New: ", "size": 11.5, "bold": True, "color": RGBColor(0xA8, 0x43, 0x1E)},
     {"text": "brightening / anti-pollution claim on the tinted variant, with real clinical "
              "substantiation indie brands can't match.", "size": 11.5}],
], line_spacing=1.35, space_after=12)

# ===========================================================================
# SLIDE: Roadmap
# ===========================================================================
s = add_slide()
header(s, "PORTFOLIO STRATEGY", "Phased rollout, Year 1 to Year 5", 7)
picture(s, f"{ASSETS}/roadmap.png", 0.6, 1.7, w=12.13)

# ===========================================================================
# SLIDE: Pack formats & pricing
# ===========================================================================
s = add_slide()
header(s, "PORTFOLIO STRATEGY", "Pack formats, sizes and pricing", 8)
picture(s, f"{ASSETS}/pricing_ladder.png", 0.5, 1.75, h=4.6)
rich(s, 8.65, 1.9, 4.1, 4.8, [
    [{"text": "Why this ladder", "size": 13, "bold": True, "color": TEAL_DARK}],
    [{"text": "Cetaphil stays mass-premium, priced against The Derma Co. and Re'equil.",
      "size": 11}],
    [{"text": "Biluma stays derma-cosmetic premium, approaching (not matching) La "
              "Roche-Posay import pricing.", "size": 11}],
    [{"text": "Trial sizes exist purely for quick-commerce discovery economics -- not a "
              "margin play.", "size": 11}],
    [{"text": "The two-tier hierarchy is preserved throughout: no SKU collapses both "
              "brands into one price band.", "size": 11}],
], line_spacing=1.35, space_after=12)

# ===========================================================================
# SLIDE: Go-to-market
# ===========================================================================
s = add_slide()
header(s, "GO-TO-MARKET", "Channels, partnerships, and the trial-to-repeat path", 9)

gtm = [
    ("CHANNELS & TERRITORIES", TEAL, [
        "Lead with quick commerce + Nykaa/Amazon for trial-size and reapplication SKUs",
        "Pharmacy & dermatologist-clinic channels reinforce the credibility claim",
        "Metro / Tier-1 first; Tier-2 expansion in Years 2-3 with mass-tier Cetaphil formats",
    ]),
    ("PARTNERSHIPS", CORAL, [
        "Dermatologist-creator partnerships, not generic beauty influencers",
        "Carries credibility into the content-first discovery journey this cohort uses",
        "A structural advantage no indie D2C brand can fully replicate",
    ]),
    ("TRIAL -> ADOPTION -> REPEAT", GOLD, [
        "Trial: small-format SKUs on quick commerce, low-commitment pricing",
        "Adoption: stick + mist are the wedge -- fix reapplication, daily use follows",
        "Repeat: subscribe/reorder mechanics, plus a Cetaphil-to-Biluma \"graduate\" path",
    ]),
]
x = 0.6
w = 3.95
for title, color, bullets in gtm:
    rect(s, x, 1.65, w, 0.55, fill=color)
    textbox(s, x + 0.15, 1.65, w - 0.3, 0.55, title, size=11.5, color=WHITE, bold=True,
            anchor=MSO_ANCHOR.MIDDLE, font=FONT)
    rect(s, x, 2.2, w, 3.9, fill=WHITE, line_color=color, line_w=1.2)
    paras = [[{"text": f"• {b}", "size": 11}] for b in bullets]
    rich(s, x + 0.2, 2.42, w - 0.4, 3.5, paras, line_spacing=1.35, space_after=14)
    x += w + 0.19

# ===========================================================================
# SLIDE: Risks
# ===========================================================================
s = add_slide()
header(s, "RISKS & CAVEATS", "What this strategy assumes, and what to validate", 10)

risks = [
    ("Reformulation risk", "\"No white cast\" is a real R&D commitment, not a marketing claim "
     "-- achievability within the phase timeline needs Galderma's own formulation science."),
    ("Cannibalization risk", "Cetaphil's new gel-cream and Biluma's existing lotion could "
     "overlap -- pricing and claims must stay differentiated by finish and skin-goal."),
    ("Distribution execution risk", "Winning quick commerce needs fulfillment and reorder "
     "readiness outside this analysis's scope -- validate with actual channel partners."),
    ("Data risk", "Market sizing and channel growth are secondary research with real variance "
     "across sources -- see research/sources.md -- treat as directional, not internal-grade."),
]
y = 1.75
for title, body in risks:
    rect(s, 0.6, y, 0.06, 1.15, fill=CORAL)
    textbox(s, 0.85, y, 3.2, 0.4, title, size=13, color=TEAL_DARK, bold=True)
    textbox(s, 0.85, y + 0.42, 11.0, 0.7, body, size=11.5, color=INK, line_spacing=1.3)
    y += 1.32

prs.save(OUT)
print("Saved:", OUT)
print("Total slides:", SLIDE_N)
